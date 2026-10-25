#
# Copyright (C) 2024-2025 EDUmind - Los Mundos Edufis
# Author: Luis Vilela Acuña
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU Affero General Public License as published by
# the Free Software Foundation, either version 3 of the License, or
# (at your option) any later version.
#
# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU Affero General Public License for more details.
#
# You should have received a copy of the GNU Affero General Public License
# along with this program.  If not, see <https://www.gnu.org/licenses/>.
#

from fastapi import APIRouter, HTTPException, Request
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from typing import Optional
from io import BytesIO
from pathlib import Path
from fpdf import FPDF
import html
import logging
import re

from app.core.rate_limit import limiter

logger = logging.getLogger(__name__)

router = APIRouter()




# Límites: la exportación es anónima y genera PDF en CPU; sin tope era un DoS trivial.
MAX_EXPORT_CHARS = 200_000


class ExportRequest(BaseModel):
    content: str = Field(max_length=MAX_EXPORT_CHARS)
    filename: str = Field(default="documento", max_length=120)
    title: Optional[str] = Field(default=None, max_length=200)


def _safe_filename(raw: str, extension: str) -> str:
    """Nombre de fichero ASCII seguro para Content-Disposition (sin comillas, saltos ni rutas)."""
    base = Path(raw or "").name
    base = re.sub(r"[^A-Za-z0-9._-]+", "_", base).strip("._-") or "documento"
    return f"{base[:80]}.{extension}"


def _content_disposition(raw: str, extension: str) -> dict:
    return {"Content-Disposition": f'attachment; filename="{_safe_filename(raw, extension)}"'}


@router.post("/pdf")
@limiter.limit("12/minute")
async def export_pdf(request: Request, payload: ExportRequest):
    """Export content to PDF format."""
    try:
        pdf = FPDF()
        pdf.add_page()
        pdf.set_auto_page_break(auto=True, margin=15)
        
        # Add title if provided
        if payload.title:
            pdf.set_font("Helvetica", "B", 16)
            pdf.multi_cell(0, 10, payload.title)
            pdf.ln(5)
        
        # Add content
        pdf.set_font("Helvetica", "", 11)
        
        # Clean and encode content for PDF
        clean_content = html.unescape(payload.content)
        # Handle Unicode by encoding to latin-1 with replacement
        try:
            pdf.multi_cell(0, 6, clean_content)
        except UnicodeEncodeError:
            # Fallback: replace non-latin characters
            safe_content = clean_content.encode('latin-1', errors='replace').decode('latin-1')
            pdf.multi_cell(0, 6, safe_content)
        
        # Output to bytes
        pdf_bytes = pdf.output()
        buffer = BytesIO(pdf_bytes)
        buffer.seek(0)
        
        return StreamingResponse(
            buffer,
            media_type="application/pdf",
            headers=_content_disposition(payload.filename, "pdf"),
        )
    except Exception:
        logger.exception("Error generando PDF")
        raise HTTPException(status_code=500, detail="No se pudo generar el PDF.")


@router.post("/txt")
@limiter.limit("12/minute")
async def export_txt(request: Request, payload: ExportRequest):
    """Export content to plain text format."""
    try:
        content = payload.content
        if payload.title:
            content = f"{payload.title}\n{'=' * len(payload.title)}\n\n{content}"
        
        buffer = BytesIO(content.encode('utf-8'))
        buffer.seek(0)
        
        return StreamingResponse(
            buffer,
            media_type="text/plain; charset=utf-8",
            headers=_content_disposition(payload.filename, "txt"),
        )
    except Exception:
        logger.exception("Error generando TXT")
        raise HTTPException(status_code=500, detail="No se pudo generar el TXT.")


@router.post("/html")
@limiter.limit("12/minute")
async def export_html(request: Request, payload: ExportRequest):
    """Export content to HTML format (can be opened in Word)."""
    try:
        # El contenido es texto del usuario: se escapa antes de insertarlo en el documento.
        html_content = html.escape(payload.content).replace('\n', '<br>\n')
        safe_title = html.escape(payload.title or payload.filename)
        
        full_html = f"""<!DOCTYPE html>
<html lang="es">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{safe_title}</title>
    <style>
        body {{
            font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif;
            max-width: 800px;
            margin: 40px auto;
            padding: 20px;
            line-height: 1.6;
            color: #333;
        }}
        h1 {{
            color: #2563eb;
            border-bottom: 2px solid #2563eb;
            padding-bottom: 10px;
        }}
    </style>
</head>
<body>
    {f'<h1>{safe_title}</h1>' if payload.title else ''}
    <div class="content">
        {html_content}
    </div>
    <footer style="margin-top: 40px; padding-top: 20px; border-top: 1px solid #eee; font-size: 12px; color: #666;">
        Generado con EDUmind AI
    </footer>
</body>
</html>"""
        
        buffer = BytesIO(full_html.encode('utf-8'))
        buffer.seek(0)
        
        return StreamingResponse(
            buffer,
            media_type="text/html; charset=utf-8",
            headers=_content_disposition(payload.filename, "html"),
        )
    except Exception:
        logger.exception("Error generando HTML")
        raise HTTPException(status_code=500, detail="No se pudo generar el HTML.")


def _build_docx(payload: ExportRequest) -> bytes:
    """Documento Word sencillo: título, encabezados (líneas cortas en mayúsculas) y viñetas (líneas con guion)."""
    from docx import Document

    documento = Document()
    if payload.title:
        documento.add_heading(payload.title, level=1)
    for linea in html.unescape(payload.content).replace("\r\n", "\n").split("\n"):
        texto = linea.rstrip()
        if not texto.strip():
            continue
        limpio = texto.strip()
        if limpio.startswith(("- ", "• ", "* ")):
            documento.add_paragraph(limpio[2:].strip(), style="List Bullet")
        elif re.match(r"^\d+[.)]\s+\S", limpio) and len(limpio) <= 120 and limpio.upper() == limpio:
            documento.add_heading(limpio, level=2)
        elif len(limpio) <= 90 and limpio.upper() == limpio and any(ch.isalpha() for ch in limpio):
            documento.add_heading(limpio.rstrip(":"), level=2)
        else:
            documento.add_paragraph(limpio)
    buffer = BytesIO()
    documento.save(buffer)
    return buffer.getvalue()


@router.post("/docx")
@limiter.limit("12/minute")
async def export_docx(request: Request, payload: ExportRequest):
    """Exporta el contenido a Word (.docx), el formato que piden los centros."""
    try:
        datos = _build_docx(payload)
        return StreamingResponse(
            BytesIO(datos),
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers=_content_disposition(payload.filename, "docx"),
        )
    except Exception:
        logger.exception("Error generando DOCX")
        raise HTTPException(status_code=500, detail="No se pudo generar el DOCX.")
